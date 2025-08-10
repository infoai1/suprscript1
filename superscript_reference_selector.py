import streamlit as st
from docx import Document
from io import BytesIO

def detect_superscripts_and_find_refs(docx_bytes):
    document = Document(BytesIO(docx_bytes))
    superscripts = []

    # Detect superscript numbers in runs
    for para in document.paragraphs:
        for run in para.runs:
            if run.font.superscript and run.text.strip().isdigit():
                superscripts.append((run.text.strip(), para.text.strip()))

    # Find paragraphs starting exactly with number + dot for each superscript
    superscript_refs = []
    for num, sup_para_text in superscripts:
        options = [para.text.strip() for para in document.paragraphs if para.text.strip().startswith(f"{num}.")]
        superscript_refs.append((num, sup_para_text, options))

    return superscript_refs

st.title("Superscript Reference Selector")

uploaded_file = st.file_uploader("Upload a DOCX file", type=["docx"])

if uploaded_file is not None:
    results = detect_superscripts_and_find_refs(uploaded_file.read())
    st.write("### Select the correct reference for each superscript")

    selections = {}
    for idx, (num, sup_text, options) in enumerate(results):
        st.markdown(f"**Superscript {num} found in paragraph:**")
        st.write(f"> {sup_text}")

        if not options:
            st.write("No matching reference paragraphs found.")
            selections[num] = None
        elif len(options) == 1:
            st.write("Only one reference found (auto-selected):")
            st.write(f"> {options[0]}")
            selections[num] = options
        else:
            # User selects one option from dropdown
            selection = st.selectbox(
                f"Choose the correct reference paragraph for superscript {num}:",
                options,
                key=f"select_{idx}"
            )
            selections[num] = selection
        st.markdown("---")

    st.write("### Your selected references:")
    for num, selected_ref in selections.items():
        if selected_ref:
            st.write(f"Superscript {num}:")
            st.write(f"> {selected_ref}")
        else:
            st.write(f"Superscript {num}: No reference selected.")
